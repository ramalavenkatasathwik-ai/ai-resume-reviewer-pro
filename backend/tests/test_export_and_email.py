import io
from datetime import datetime, timezone
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
import docx

from app.core.config import get_settings
from app.core.security import create_access_token, hash_password
from app.models.user import User
from app.models.resume import ResumeDocument, ResumeVersion
from app.models.analysis import AnalysisHistory
from app.models.job_match import JobMatchReport
from app.models.email_log import EmailDeliveryLog
from app.services.pdf_export_service import PdfExportService
from app.services.docx_export_service import DocxExportService
from app.services.email_service import EmailService


@pytest.fixture
def export_user(test_db_session: Session) -> User:
    user = User(
        email="export_tester@example.com",
        hashed_password=hash_password("SecurePassword123!"),
        display_name="Export Tester",
    )
    test_db_session.add(user)
    test_db_session.commit()
    test_db_session.refresh(user)
    return user


@pytest.fixture
def export_other_user(test_db_session: Session) -> User:
    user = User(
        email="other_export@example.com",
        hashed_password=hash_password("SecurePassword123!"),
        display_name="Other User",
    )
    test_db_session.add(user)
    test_db_session.commit()
    test_db_session.refresh(user)
    return user


@pytest.fixture
def export_auth_headers(export_user: User) -> dict:
    token = create_access_token({"sub": export_user.id})
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def other_export_auth_headers(export_other_user: User) -> dict:
    token = create_access_token({"sub": export_other_user.id})
    return {"Authorization": f"Bearer {token}"}


SAMPLE_RICH_RESUME = {
    "basics": {
        "name": "Jane Doe",
        "email": "jane.doe@example.com",
        "phone": "+1 555-0199",
        "location": "San Francisco, CA",
        "website": "https://janedoe.dev",
        "linkedin": "https://linkedin.com/in/janedoe",
        "github": "https://github.com/janedoe",
        "title": "Lead AI Platform Engineer",
        "summary": "Proven AI systems engineer with 8+ years designing scalable LLM inference pipelines, RAG engines, and Kubernetes microservices.",
    },
    "experience": [
        {
            "company": "DeepTech Systems Inc.",
            "position": "Principal Architect",
            "location": "San Francisco, CA",
            "startDate": "2022-03",
            "endDate": "Present",
            "highlights": [
                "Led team of 12 engineers building multi-tenant agent execution framework handling 5M daily queries.",
                "Optimized inference throughput by 3.2x while reducing GPU cluster costs by $240,000 annually.",
                "Architected streaming token delivery engine with p99 response latency under 120ms.",
            ],
        },
        {
            "company": "Cognitive Cloud Labs",
            "position": "Senior Backend Engineer",
            "location": "Mountain View, CA",
            "startDate": "2019-01",
            "endDate": "2022-02",
            "highlights": [
                "Designed distributed indexing pipeline over PostgreSQL and Qdrant vector databases.",
                "Built automated CI/CD and deployment workflows on AWS EKS with zero-downtime rollouts.",
            ],
        },
    ],
    "education": [
        {
            "institution": "University of California, Berkeley",
            "degree": "B.S. in Computer Science & Applied Mathematics",
            "location": "Berkeley, CA",
            "startDate": "2014-08",
            "endDate": "2018-05",
            "highlights": ["Graduated Magna Cum Laude", "Head TA for Distributed Computing"],
        }
    ],
    "skills": [
        "Python",
        "FastAPI",
        "PyTorch",
        "PostgreSQL",
        "Docker",
        "Kubernetes",
        "Vector Databases",
        "TypeScript",
        "Next.js",
        "Distributed Systems",
    ],
    "projects": [
        {
            "name": "OpenAgent Stream",
            "description": "High-throughput asynchronous LLM pipeline framework with native telemetry.",
            "url": "https://github.com/janedoe/openagent",
            "highlights": ["Over 3,500 GitHub stars and 40+ active open-source contributors."],
        }
    ],
    "certifications": ["AWS Certified Solutions Architect - Professional", "CKS: Certified Kubernetes Security Specialist"],
}


# ==============================================================================
# 1. Direct Service Unit Tests (PDF & DOCX)
# ==============================================================================


def test_pdf_export_service_ats_classic():
    pdf_bytes = PdfExportService.generate_resume_pdf(
        resume_data=SAMPLE_RICH_RESUME,
        template_name="ats_classic",
        page_format="letter",
    )
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 1000
    assert pdf_bytes.startswith(b"%PDF-")


def test_pdf_export_service_modern_professional_a4():
    pdf_bytes = PdfExportService.generate_resume_pdf(
        resume_data=SAMPLE_RICH_RESUME,
        template_name="modern",
        page_format="a4",
    )
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 1000
    assert pdf_bytes.startswith(b"%PDF-")


def test_docx_export_service_resume_generation():
    docx_bytes = DocxExportService.generate_resume_docx(
        resume_data=SAMPLE_RICH_RESUME,
        template_name="modern",
        page_format="letter",
    )
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 2000

    # Parse with python-docx to verify document structure
    doc = docx.Document(io.BytesIO(docx_bytes))
    all_text = " ".join([p.text for p in doc.paragraphs])
    assert "Jane Doe" in all_text
    assert "Lead AI Platform Engineer" in all_text
    assert "DeepTech Systems Inc." in all_text
    assert "Principal Architect" in all_text
    assert "Python" in all_text or any("Python" in cell.text for t in doc.tables for row in t.rows for cell in row.cells)


def test_analysis_report_pdf_generation(mock_analysis_dict: dict):
    pdf_bytes = PdfExportService.generate_analysis_report_pdf(
        report_id="test-analysis-123",
        target_role="Senior AI Engineer",
        original_filename="sample_resume.pdf",
        created_at=datetime.now(timezone.utc),
        overall_score=88,
        ats_score=92,
        result_json=mock_analysis_dict,
        source_type="pdf",
    )
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 1000
    assert pdf_bytes.startswith(b"%PDF-")


def test_job_match_report_pdf_generation():
    comparison_items = [
        {
            "category": "required_skills",
            "requirement": "5+ years Python experience",
            "importance": "high",
            "match_status": "matched",
            "resume_evidence": "Principal Architect with 8+ years building Python LLM pipelines.",
            "recommendation": None,
        },
        {
            "category": "required_skills",
            "requirement": "Apache Kafka event streaming",
            "importance": "high",
            "match_status": "missing",
            "resume_evidence": None,
            "recommendation": "Explicitly highlight pub-sub or Kafka streaming projects.",
        },
    ]

    pdf_bytes = PdfExportService.generate_job_match_report_pdf(
        report_id="job-match-456",
        job_title="Lead Platform Engineer",
        company_name="Anthropic AI",
        overall_score=84,
        required_skills_score=80,
        preferred_skills_score=90,
        experience_score=85,
        keyword_score=82,
        comparison_items=comparison_items,
        recommendations=["Highlight Kafka event streaming", "Quantify cluster latency metrics"],
        matched_keywords=["Python", "FastAPI", "Docker", "Kubernetes"],
        missing_keywords=["Kafka", "Terraform"],
        created_at=datetime.now(timezone.utc),
    )
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 1000
    assert pdf_bytes.startswith(b"%PDF-")


# ==============================================================================
# 2. Export API Endpoints (Auth, Formats, Ownership)
# ==============================================================================


@pytest.fixture
def saved_resume_and_version(test_db_session: Session, export_user: User):
    resume = ResumeDocument(
        user_id=export_user.id,
        title="AI Engineer Master",
        target_role="Senior AI Engineer",
    )
    test_db_session.add(resume)
    test_db_session.commit()
    test_db_session.refresh(resume)

    version = ResumeVersion(
        resume_id=resume.id,
        user_id=export_user.id,
        version_number=1,
        version_notes="Initial baseline",
        resume_data=SAMPLE_RICH_RESUME,
        template_name="modern",
        page_format="letter",
    )
    test_db_session.add(version)
    test_db_session.commit()
    test_db_session.refresh(version)

    resume.current_version_id = version.id
    test_db_session.commit()
    test_db_session.refresh(resume)

    return resume, version


def test_export_resume_version_pdf_success(
    client: TestClient,
    export_auth_headers: dict,
    saved_resume_and_version: tuple[ResumeDocument, ResumeVersion],
):
    resume, version = saved_resume_and_version
    res = client.get(
        f"/api/resumes/{resume.id}/versions/{version.id}/export/pdf",
        headers=export_auth_headers,
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert "attachment;" in res.headers["content-disposition"]
    assert res.content.startswith(b"%PDF-")


def test_export_resume_version_docx_success(
    client: TestClient,
    export_auth_headers: dict,
    saved_resume_and_version: tuple[ResumeDocument, ResumeVersion],
):
    resume, version = saved_resume_and_version
    res = client.get(
        f"/api/resumes/{resume.id}/versions/{version.id}/export/docx",
        headers=export_auth_headers,
    )
    assert res.status_code == 200
    assert "officedocument" in res.headers["content-type"]
    assert "attachment;" in res.headers["content-disposition"]
    assert len(res.content) > 1000


def test_export_resume_cross_user_forbidden(
    client: TestClient,
    other_export_auth_headers: dict,
    saved_resume_and_version: tuple[ResumeDocument, ResumeVersion],
):
    resume, version = saved_resume_and_version
    res = client.get(
        f"/api/resumes/{resume.id}/versions/{version.id}/export/pdf",
        headers=other_export_auth_headers,
    )
    assert res.status_code in (403, 404)


def test_export_unsaved_preview_pdf_and_docx(
    client: TestClient,
    export_auth_headers: dict,
    saved_resume_and_version: tuple[ResumeDocument, ResumeVersion],
):
    resume, _ = saved_resume_and_version
    payload = {
        "resume_data": SAMPLE_RICH_RESUME,
        "template_name": "ats_classic",
        "page_format": "letter",
    }

    # PDF preview
    pdf_res = client.post(
        f"/api/resumes/{resume.id}/export/pdf-preview",
        json=payload,
        headers=export_auth_headers,
    )
    assert pdf_res.status_code == 200
    assert pdf_res.headers["content-type"] == "application/pdf"
    assert pdf_res.content.startswith(b"%PDF-")

    # DOCX preview
    docx_res = client.post(
        f"/api/resumes/{resume.id}/export/docx-preview",
        json=payload,
        headers=export_auth_headers,
    )
    assert docx_res.status_code == 200
    assert "officedocument" in docx_res.headers["content-type"]
    assert len(docx_res.content) > 1000


def test_export_analysis_report_pdf_and_authorization(
    client: TestClient,
    test_db_session: Session,
    export_user: User,
    export_auth_headers: dict,
    other_export_auth_headers: dict,
    mock_analysis_dict: dict,
):
    history = AnalysisHistory(
        user_id=export_user.id,
        target_role="Senior Platform Engineer",
        original_filename="jane_cv.pdf",
        overall_score=88,
        ats_score=92,
        result_json=mock_analysis_dict,
        source_type="pdf",
    )
    test_db_session.add(history)
    test_db_session.commit()
    test_db_session.refresh(history)

    # Owner can download
    res = client.get(
        f"/api/analyses/{history.id}/export/pdf",
        headers=export_auth_headers,
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF-")

    # Other user denied (403)
    res_other = client.get(
        f"/api/analyses/{history.id}/export/pdf",
        headers=other_export_auth_headers,
    )
    assert res_other.status_code == 403


def test_export_job_match_report_pdf_and_authorization(
    client: TestClient,
    test_db_session: Session,
    export_user: User,
    export_auth_headers: dict,
    other_export_auth_headers: dict,
    saved_resume_and_version: tuple[ResumeDocument, ResumeVersion],
):
    resume, version = saved_resume_and_version
    job_match = JobMatchReport(
        user_id=export_user.id,
        resume_id=resume.id,
        resume_version_id=version.id,
        job_title="Lead AI Engineer",
        company_name="DeepMind",
        job_description_raw="Seeking a lead engineer proficient in PyTorch, distributed training and Python.",
        extracted_requirements={"required_skills": ["PyTorch", "Python"]},
        comparison_result={"items": [], "score_breakdown": {"matched_keywords": ["PyTorch"]}},
        recommendations=["Add more distributed training metrics."],
        overall_score=85,
        required_skills_score=85,
        preferred_skills_score=80,
        experience_score=90,
        keyword_score=85,
    )
    test_db_session.add(job_match)
    test_db_session.commit()
    test_db_session.refresh(job_match)

    # Owner can download
    res = client.get(
        f"/api/job-matches/{job_match.id}/export/pdf",
        headers=export_auth_headers,
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF-")

    # Other user cannot access
    res_other = client.get(
        f"/api/job-matches/{job_match.id}/export/pdf",
        headers=other_export_auth_headers,
    )
    assert res_other.status_code in (403, 404)


# ==============================================================================
# 3. Email Delivery API Tests (Status, Unconfigured, Validation, Duplicate, Mocked Send)
# ==============================================================================


def test_email_status_unconfigured(client: TestClient, export_auth_headers: dict, monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "EMAIL_PROVIDER", None)
    monkeypatch.setattr(settings, "SMTP_HOST", None)
    monkeypatch.setattr(settings, "EMAIL_API_KEY", None)

    res = client.get("/api/email/status", headers=export_auth_headers)
    assert res.status_code == 200
    body = res.json()
    assert body["is_configured"] is False
    assert body["from_address"] is None


def test_email_send_when_unconfigured_returns_503(
    client: TestClient,
    export_auth_headers: dict,
    saved_resume_and_version: tuple[ResumeDocument, ResumeVersion],
    monkeypatch,
):
    settings = get_settings()
    monkeypatch.setattr(settings, "EMAIL_PROVIDER", None)
    monkeypatch.setattr(settings, "SMTP_HOST", None)
    monkeypatch.setattr(settings, "EMAIL_API_KEY", None)

    resume, version = saved_resume_and_version
    payload = {
        "recipient_email": "recruiter@example.com",
        "document_type": "resume_pdf",
        "document_id": resume.id,
        "version_id": version.id,
        "template_name": "modern",
        "page_format": "letter",
    }
    res = client.post("/api/email/send", json=payload, headers=export_auth_headers)
    assert res.status_code == 503
    assert res.json()["detail"]["error"] == "EMAIL_NOT_CONFIGURED"


def test_email_send_validation_and_header_injection_rejection(
    client: TestClient,
    export_auth_headers: dict,
    saved_resume_and_version: tuple[ResumeDocument, ResumeVersion],
    monkeypatch,
):
    settings = get_settings()
    monkeypatch.setattr(settings, "EMAIL_PROVIDER", "smtp")
    monkeypatch.setattr(settings, "SMTP_HOST", "smtp.example.com")
    monkeypatch.setattr(settings, "EMAIL_FROM_ADDRESS", "noreply@example.com")

    resume, version = saved_resume_and_version

    # Invalid email address format
    bad_email_payload = {
        "recipient_email": "not-an-email",
        "document_type": "resume_pdf",
        "document_id": resume.id,
        "version_id": version.id,
    }
    res = client.post("/api/email/send", json=bad_email_payload, headers=export_auth_headers)
    assert res.status_code in (400, 422)

    # CRLF header injection in custom subject
    injection_payload = {
        "recipient_email": "legit@example.com",
        "document_type": "resume_pdf",
        "document_id": resume.id,
        "version_id": version.id,
        "subject": "Legit Subject\r\nBcc: evil@attacker.com",
    }
    res_inject = client.post("/api/email/send", json=injection_payload, headers=export_auth_headers)
    assert res_inject.status_code in (400, 422)


@pytest.mark.asyncio
async def test_email_send_success_mocked_and_duplicate_protection(
    client: TestClient,
    test_db_session: Session,
    export_user: User,
    export_auth_headers: dict,
    saved_resume_and_version: tuple[ResumeDocument, ResumeVersion],
    monkeypatch,
):
    settings = get_settings()
    monkeypatch.setattr(settings, "EMAIL_PROVIDER", "smtp")
    monkeypatch.setattr(settings, "SMTP_HOST", "smtp.example.com")
    monkeypatch.setattr(settings, "SMTP_PORT", 587)
    monkeypatch.setattr(settings, "EMAIL_FROM_ADDRESS", "noreply@resumereviewer.pro")
    monkeypatch.setattr(settings, "EMAIL_FROM_NAME", "AI Resume Reviewer Pro")

    # Mock the internal low-level SMTP send to succeed and return a mock message ID
    async def mock_send_smtp(*args, **kwargs):
        return "mock-smtp-msgid-998877"

    monkeypatch.setattr(EmailService, "_send_via_smtp", mock_send_smtp)

    resume, version = saved_resume_and_version
    payload = {
        "recipient_email": "hiring.manager@example.com",
        "document_type": "resume_pdf",
        "document_id": resume.id,
        "version_id": version.id,
        "template_name": "modern",
        "page_format": "letter",
        "subject": "Application for Lead AI Engineer",
        "custom_message": "Please find attached my updated resume.",
    }

    # 1. First send -> Success
    res = client.post("/api/email/send", json=payload, headers=export_auth_headers)
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "accepted"
    assert data["recipient_email"] == "hiring.manager@example.com"
    assert data["provider"] == "smtp"

    # Verify log entry in database
    log = (
        test_db_session.query(EmailDeliveryLog)
        .filter(EmailDeliveryLog.recipient_email == "hiring.manager@example.com")
        .first()
    )
    assert log is not None
    assert log.status == "accepted"
    assert log.document_type == "resume_pdf"
    assert log.user_id == export_user.id
    assert "mock-smtp-msgid" in (log.provider_message_id or "")

    # 2. Duplicate send attempt within 60s cooldown -> HTTP 429
    res_dup = client.post("/api/email/send", json=payload, headers=export_auth_headers)
    assert res_dup.status_code == 429
    assert res_dup.json()["detail"]["error"] == "DUPLICATE_SEND_BLOCKED"
